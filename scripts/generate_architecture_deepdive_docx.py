# -*- coding: utf-8 -*-
"""
Generator for:
1. Thuyet_Trinh_Chuyen_Sau_Kien_Truc_Du_An.docx
2. THUYET_TRINH_KIEN_TRUC_VA_DU_LIEU.md

Comprehensive presentation and viva defense document on Core Architecture, Data Storage,
State Lifecycle, and Technical Trade-offs for Study Planner (DNTU - GV: Thầy Nguyễn Xuân Vinh).
"""

import os
import sys
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from architecture_deepdive_data import PROJECT_METADATA, ARCHITECTURAL_ISSUES

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_left_border(cell, color_hex="2563EB", size="36"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:left w:val="single" w:sz="{size}" w:space="0" w:color="{color_hex}"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def set_table_borders(table, border_color="CBD5E1"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="{border_color}"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138) # Navy Blue
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(37, 99, 235) # Royal Blue
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
    return p

def add_paragraph(doc, text="", space_before=2, space_after=4, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, italic=False, color_rgb=None, size=10):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        if color_rgb:
            r.font.color.rgb = color_rgb
        else:
            r.font.color.rgb = RGBColor(30, 41, 59)
    return p

def create_callout_box(doc, text_runs, border_color="2563EB", bg_color="F8FAFC"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.8)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    set_cell_left_border(cell, color_hex=border_color, size="36")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    for r_text, r_bold, r_italic, r_color, r_size in text_runs:
        r = p.add_run(r_text)
        r.font.name = "Times New Roman"
        r.font.bold = r_bold
        r.font.italic = r_italic
        if r_color:
            r.font.color.rgb = r_color
        if r_size:
            r.font.size = Pt(r_size)
        else:
            r.font.size = Pt(9.5)
            
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(2)

def add_issue_section(doc, issue):
    """Render a comprehensive issue section in the Word document."""
    add_heading_2(doc, f"VẤN ĐỀ {issue['num']:02d}: {issue['title']}")
    
    # Meta table
    tbl = doc.add_table(rows=3, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl, "E2E8F0")
    
    col_widths = [Inches(1.8), Inches(5.0)]
    meta_rows = [
        ("Lĩnh vực chuyên sâu:", issue['category']),
        ("Thành viên phụ trách:", issue['owner']),
        ("Tiêu chuẩn đồ án đáp ứng:", issue['requirement'])
    ]
    
    for i, (k, v) in enumerate(meta_rows):
        row = tbl.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = col_widths[0], col_widths[1]
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF" if i % 2 == 0 else "F8FAFC")
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.name = "Times New Roman"
        r0.font.bold = True
        r0.font.size = Pt(9)
        r0.font.color.rgb = RGBColor(71, 85, 105)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(9)
        r1.font.color.rgb = RGBColor(15, 23, 42)
        if "TV1" in v or "Trâm" in v:
            r1.font.bold = True
        elif "TV2" in v or "Giang" in v:
            r1.font.bold = True
            
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(2)
    p_sp.paragraph_format.space_after = Pt(2)

    # 1. Problem Statement
    add_heading_3(doc, "1. Vấn Đề Đặt Ra & Nguy Cơ Tiềm Ẩn:")
    add_paragraph(doc, issue['problem_statement'], space_before=1, space_after=3)

    # 2. Storage Location & Technology
    add_heading_3(doc, "2. Dữ Liệu Được Lưu Ở Đâu & Công Nghệ Áp Dụng:")
    add_paragraph(doc, issue['storage_location'], space_before=1, space_after=3, bold=False, color_rgb=RGBColor(30, 58, 138))

    # 3. Why Not Alternative
    add_heading_3(doc, "3. Phản Biện Giải Pháp & Tại Sao Không Chọn Công Nghệ Khác:")
    add_paragraph(doc, issue['why_not_alternative'], space_before=1, space_after=3, italic=False, color_rgb=RGBColor(51, 65, 85))

    # 4. Technical Mechanism
    if issue.get('technical_mechanism'):
        add_heading_3(doc, "4. Cơ Chế Hoạt Động & Điểm Nhấn Mã Nguồn:")
        for sym_name, sym_desc in issue['technical_mechanism']:
            p_sym = doc.add_paragraph()
            p_sym.paragraph_format.space_before = Pt(1)
            p_sym.paragraph_format.space_after = Pt(1)
            p_sym.paragraph_format.left_indent = Inches(0.2)
            
            r_bullet = p_sym.add_run("• ")
            r_bullet.font.bold = True
            r_bullet.font.color.rgb = RGBColor(37, 99, 235)
            
            r_sym = p_sym.add_run(f"[{sym_name}]: ")
            r_sym.font.name = "Times New Roman"
            r_sym.font.bold = True
            r_sym.font.size = Pt(9.5)
            r_sym.font.color.rgb = RGBColor(30, 58, 138)
            
            r_desc = p_sym.add_run(sym_desc)
            r_desc.font.name = "Times New Roman"
            r_desc.font.size = Pt(9.5)
            r_desc.font.color.rgb = RGBColor(51, 65, 85)

    # 5. Presentation Speech
    if issue.get('speech'):
        create_callout_box(
            doc,
            [
                ("🗣️ LỜI THOẠI THUYẾT TRÌNH MẪU KHI TRÌNH BÀY VẤN ĐỀ NÀY:\n", True, False, RGBColor(5, 150, 105), 10),
                (f'"{issue["speech"]}"', False, True, RGBColor(15, 23, 42), 9.5)
            ],
            border_color="059669", # Emerald Green
            bg_color="F0FDF4"
        )

    # 6. Viva Q&A
    if issue.get('viva_qna'):
        trap = issue['viva_qna']
        create_callout_box(
            doc,
            [
                ("🎯 BẪY VẤN ĐÁP CỦA THẦY NGUYỄN XUÂN VINH & CÁCH TRẢ LỜI ĐẠT ĐIỂM 10:\n", True, False, RGBColor(180, 83, 9), 10),
                ("❓ Thầy hỏi: ", True, False, RGBColor(180, 83, 9), 9.5),
                (f'"{trap["question"]}"\n', True, True, RGBColor(180, 83, 9), 9.5),
                ("💡 Trả lời tự tin: ", True, False, RGBColor(30, 58, 138), 9.5),
                (f'"{trap["answer"]}"', False, False, RGBColor(15, 23, 42), 9.5)
            ],
            border_color="D97706", # Amber
            bg_color="FFFBEB"
        )
    
    # Separator
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(4)
    p_div.paragraph_format.space_after = Pt(6)
    r_div = p_div.add_run("· " * 42)
    r_div.font.color.rgb = RGBColor(203, 213, 225)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER

def build_architecture_documents(output_docx_path, output_md_path):
    doc = docx.Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    # Title & Metadata
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_after = Pt(2)
    r_inst = p_inst.add_run(PROJECT_METADATA["university"].upper() + "\nBỘ MÔN CÔNG NGHỆ PHẦN MỀM & KỸ THUẬT DỮ LIỆU")
    r_inst.font.name = "Times New Roman"
    r_inst.font.size = Pt(10)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(100, 116, 139)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("CHUYÊN ĐỀ THUYẾT MINH KIẾN TRÚC & LƯU TRỮ DỮ LIỆU\n(DEEP-DIVE ARCHITECTURAL SPECIFICATION & DATA LIFECYCLE)")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(15)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 58, 138)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(6)
    r_sub = p_sub.add_run(
        f"{PROJECT_METADATA['topic']} | Môn: {PROJECT_METADATA['course']}\n"
        f"Giảng viên hướng dẫn & chấm thi: {PROJECT_METADATA['instructor']}\n"
        f"Sinh viên thực hiện: {PROJECT_METADATA['member1']}\n"
        f"                   {PROJECT_METADATA['member2']}\n"
        f"Phiên bản: {PROJECT_METADATA['version']} | Kho mã nguồn Git: {PROJECT_METADATA['repo']}"
    )
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(9.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(71, 85, 105)
    
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(8)
    r_div = p_div.add_run("═" * 56)
    r_div.font.color.rgb = RGBColor(148, 163, 184)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # SECTION: SUMMARY MATRIX
    add_heading_1(doc, "BẢNG TỔNG QUAN MA TRẬN 12 VẤN ĐỀ KỸ THUẬT & NƠI LƯU TRỮ DỮ LIỆU")
    add_paragraph(
        doc,
        "Dưới đây là ma trận tóm tắt toàn bộ 12 bài toán kiến trúc cốt lõi trong ứng dụng Study Planner, "
        "đối chiếu trực tiếp với vị trí lưu trữ dữ liệu và giải pháp công nghệ đã triển khai:",
        space_before=2, space_after=4
    )
    
    tbl_sum = doc.add_table(rows=len(ARCHITECTURAL_ISSUES) + 1, cols=4)
    tbl_sum.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_sum.autofit = False
    set_table_borders(tbl_sum, "CBD5E1")
    
    s_widths = [Inches(0.6), Inches(2.3), Inches(2.2), Inches(1.7)]
    s_headers = ["STT", "Vấn đề kỹ thuật đặt ra", "Nơi lưu trữ / Công nghệ", "Thành viên phụ trách"]
    
    for j, h in enumerate(s_headers):
        cell = tbl_sum.rows[0].cells[j]
        cell.width = s_widths[j]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 80, 80, 80, 80)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Times New Roman"
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    for i, issue in enumerate(ARCHITECTURAL_ISSUES):
        row = tbl_sum.rows[i + 1]
        c0, c1, c2, c3 = row.cells[0], row.cells[1], row.cells[2], row.cells[3]
        c0.width, c1.width, c2.width, c3.width = s_widths[0], s_widths[1], s_widths[2], s_widths[3]
        for c in [c0, c1, c2, c3]:
            set_cell_margins(c, 50, 50, 80, 80)
            set_cell_background(c, "F8FAFC" if i % 2 == 1 else "FFFFFF")
            
        r0 = c0.paragraphs[0].add_run(f"{issue['num']:02d}")
        r0.font.name = "Times New Roman"; r0.font.size = Pt(8.5); r0.font.bold = True
        
        r1 = c1.paragraphs[0].add_run(issue['title'])
        r1.font.name = "Times New Roman"; r1.font.size = Pt(8.5); r1.font.bold = True
        r1.font.color.rgb = RGBColor(30, 58, 138)
        
        # Extract short storage location
        short_loc = issue['storage_location'].split("\n")[0]
        r2 = c2.paragraphs[0].add_run(short_loc)
        r2.font.name = "Times New Roman"; r2.font.size = Pt(8.5)
        
        r3 = c3.paragraphs[0].add_run(issue['owner'].split(" - ")[0] if " - " in issue['owner'] else issue['owner'])
        r3.font.name = "Times New Roman"; r3.font.size = Pt(8.5); r3.font.bold = True
        if "TV1" in issue['owner'] or "Trâm" in issue['owner']:
            r3.font.color.rgb = RGBColor(5, 150, 105)
        elif "TV2" in issue['owner'] or "Giang" in issue['owner']:
            r3.font.color.rgb = RGBColor(37, 99, 235)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(6)
    p_sp.paragraph_format.space_after = Pt(6)

    # DETAILED ISSUES RENDER
    add_heading_1(doc, "THUYẾT MINH CHI TIẾT 12 VẤN ĐỀ KIẾN TRÚC & LƯU TRỮ DỮ LIỆU")
    
    for issue in ARCHITECTURAL_ISSUES:
        add_issue_section(doc, issue)

    # FINAL SECTION: GOLDEN VIVA CHEAT SHEET
    add_heading_1(doc, "BẢNG VÀNG VẤN ĐÁP KỸ THUẬT (GOLDEN VIVA CHEAT-SHEET)")
    add_paragraph(
        doc,
        "Tóm tắt 1 câu duy nhất cho từng vấn đề để sinh viên trả lời ngay lập tức trong 3 giây đầu tiên khi Thầy Vinh đặt câu hỏi:",
        space_before=2, space_after=4
    )
    
    tbl_gold = doc.add_table(rows=len(ARCHITECTURAL_ISSUES) + 1, cols=3)
    tbl_gold.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_gold.autofit = False
    set_table_borders(tbl_gold, "CBD5E1")
    
    g_widths = [Inches(2.0), Inches(1.3), Inches(3.5)]
    g_headers = ["Câu hỏi của Thầy", "Thành viên", "Câu trả lời vàng (3 giây đầu ghi điểm 10)"]
    
    for j, h in enumerate(g_headers):
        cell = tbl_gold.rows[0].cells[j]
        cell.width = g_widths[j]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 80, 80, 80, 80)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Times New Roman"
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    golden_answers = [
        ("Token đăng nhập lưu ở đâu?", "TV2 (Giang)", "'Lưu trong chip mã hóa phần cứng SecureStore (Keychain/Keystore), tuyệt đối cấm AsyncStorage.'"),
        ("Dữ liệu bài tập lưu ở đâu?", "TV2 (Giang)", "'Lưu trong bộ nhớ đệm RAM của TanStack Query v5 với staleTime 1 phút và retry 2 lần.'"),
        ("Mất mạng dữ liệu lưu ở đâu?", "Cả 2", "'Lưu trong AsyncStorage kèm mốc thời gian thực tế hiển thị trên Banner hổ phách.'"),
        ("Optimistic Update hoạt động ra sao?", "TV2 (Giang)", "'onMutate chụp snapshot cache cũ -> onError rollback -> onSettled đồng bộ server.'"),
        ("Tiến độ % lưu ở đâu trong DB?", "TV2 (Giang)", "'Không lưu trong DB! Là Derived State O(N) tính toán động để tránh Data Inconsistency.'"),
        ("Ảnh đính kèm lưu ở đâu?", "TV1 (Trâm)", "'Lưu đường dẫn Local URI file:/// trong sandbox app, không lưu chuỗi Base64 tránh nghẽn RAM.'"),
        ("Tại sao không lỗi múi giờ 48h?", "TV2 (Giang)", "'Dùng hàm startOfDay quy chuẩn về 00:00:00 ngày lịch, không cộng trừ mili-giây chay.'"),
        ("Deep Link hoạt động thế nào?", "TV1 (Trâm)", "'Payload notification chứa assignmentId, Listener ở Root Layout đón nhận điều hướng sang [id].'"),
        ("Luật Zero-Hex hoạt động ra sao?", "TV1 (Trâm)", "'Mã hex tập trung duy nhất ở tokens.ts, UI dùng Semantic Tokens đổi Dark Mode tự động.'"),
        ("Vùng chạm 44x44pt ở đâu?", "TV1 (Trâm)", "'Cố định minHeight/minWidth 44pt tại AccessibleButton và ô checkbox, đạt chuẩn Apple HIG.'"),
        ("StateHandler làm nhiệm vụ gì?", "TV2 (Giang)", "'Xử lý đủ 4 trạng thái: Loading, Empty, Error có nút Retry gọi refetch thật, và Content.'"),
        ("Tại sao phân tách 4 vùng state?", "Cả 2", "'Áp dụng Separation of Concerns: Server State (TanStack), Client State (Zustand), Offline (AsyncStorage), Secure (SecureStore).'")
    ]
    
    for i, (q, m, a) in enumerate(golden_answers):
        row = tbl_gold.rows[i + 1]
        c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
        c0.width, c1.width, c2.width = g_widths[0], g_widths[1], g_widths[2]
        for c in [c0, c1, c2]:
            set_cell_margins(c, 50, 50, 80, 80)
            set_cell_background(c, "F8FAFC" if i % 2 == 1 else "FFFFFF")
            
        r0 = c0.paragraphs[0].add_run(q)
        r0.font.name = "Times New Roman"; r0.font.size = Pt(8.5); r0.font.bold = True
        r0.font.color.rgb = RGBColor(30, 58, 138)
        
        r1 = c1.paragraphs[0].add_run(m)
        r1.font.name = "Times New Roman"; r1.font.size = Pt(8.5); r1.font.bold = True
        r1.font.color.rgb = RGBColor(5, 150, 105) if "Trâm" in m else RGBColor(37, 99, 235)
        
        r2 = c2.paragraphs[0].add_run(a)
        r2.font.name = "Times New Roman"; r2.font.size = Pt(8.5); r2.font.italic = True
        r2.font.color.rgb = RGBColor(15, 23, 42)

    doc.save(output_docx_path)
    print(f"Word document saved to: {output_docx_path}")
    
    # BUILD COMPREHENSIVE MARKDOWN DOCUMENT
    md_content = build_markdown_document()
    with open(output_md_path, "w", encoding="utf-8") as f:
        f.write(md_content)
    print(f"Markdown document saved to: {output_md_path}")

def build_markdown_document():
    lines = []
    lines.append(f"# {PROJECT_METADATA['topic'].upper()}")
    lines.append("## CHUYÊN ĐỀ THUYẾT MINH KIẾN TRÚC & LƯU TRỮ DỮ LIỆU")
    lines.append("### (DEEP-DIVE ARCHITECTURAL SPECIFICATION & DATA LIFECYCLE)")
    lines.append("")
    lines.append(f"> **Môn học**: {PROJECT_METADATA['course']}  ")
    lines.append(f"> **Đơn vị đào tạo**: {PROJECT_METADATA['university']}  ")
    lines.append(f"> **Giảng viên hướng dẫn & chấm thi**: **{PROJECT_METADATA['instructor']}**  ")
    lines.append(f"> **Sinh viên thực hiện**:  ")
    lines.append(f"> - **{PROJECT_METADATA['member1']}**  ")
    lines.append(f"> - **{PROJECT_METADATA['member2']}**  ")
    lines.append(f"> **Phiên bản**: {PROJECT_METADATA['version']} | **Kho mã nguồn Git**: [{PROJECT_METADATA['repo']}]({PROJECT_METADATA['repo']})  ")
    lines.append("")
    lines.append("---")
    lines.append("")
    lines.append("## 1. BẢNG MA TRẬN 12 VẤN ĐỀ KỸ THUẬT & NƠI LƯU TRỮ DỮ LIỆU")
    lines.append("")
    lines.append("| STT | Vấn đề kỹ thuật đặt ra | Nơi lưu trữ / Công nghệ áp dụng | Phụ trách |")
    lines.append("| :---: | :--- | :--- | :---: |")
    for issue in ARCHITECTURAL_ISSUES:
        short_loc = issue['storage_location'].split("\n")[0]
        lines.append(f"| {issue['num']:02d} | **{issue['title']}** | {short_loc} | {issue['owner'].split(' - ')[0]} |")
    lines.append("")
    lines.append("---")
    lines.append("")
    lines.append("## 2. THUYẾT MINH CHI TIẾT 12 VẤN ĐỀ KIẾN TRÚC CỐT LÕI")
    lines.append("")
    
    for issue in ARCHITECTURAL_ISSUES:
        lines.append(f"### 📌 VẤN ĐỀ {issue['num']:02d}: {issue['title']}")
        lines.append(f"- **Lĩnh vực chuyên sâu**: {issue['category']}")
        lines.append(f"- **Thành viên phụ trách**: **{issue['owner']}**")
        lines.append(f"- **Tiêu chuẩn đồ án đáp ứng**: {issue['requirement']}")
        lines.append("")
        lines.append("#### 1. Vấn Đề Đặt Ra & Nguy Cơ Tiềm Ẩn:")
        lines.append(issue['problem_statement'])
        lines.append("")
        lines.append("#### 2. Dữ Liệu Được Lưu Ở Đâu & Công Nghệ Áp Dụng:")
        lines.append(issue['storage_location'])
        lines.append("")
        lines.append("#### 3. Phản Biện Giải Pháp & Tại Sao Không Chọn Công Nghệ Khác:")
        lines.append(issue['why_not_alternative'])
        lines.append("")
        lines.append("#### 4. Cơ Chế Hoạt Động & Điểm Nhấn Mã Nguồn:")
        for sym, desc in issue.get('technical_mechanism', []):
            lines.append(f"- **`[{sym}]`**: {desc}")
        lines.append("")
        if issue.get('speech'):
            lines.append("> **🗣️ Lời thoại thuyết trình mẫu khi trình bày vấn đề này:**")
            lines.append(f"> *\"{issue['speech']}\"*")
            lines.append("")
        if issue.get('viva_qna'):
            trap = issue['viva_qna']
            lines.append("> **🎯 Bẫy vấn đáp của Thầy Nguyễn Xuân Vinh & Cách trả lời đạt điểm 10:**")
            lines.append(f"> - **Thầy hỏi**: *\"{trap['question']}\"*")
            lines.append(f"> - **Sinh viên trả lời**: *\"{trap['answer']}\"*")
            lines.append("")
        lines.append("---")
        lines.append("")

    lines.append("## 3. BẢNG VÀNG VẤN ĐÁP KỸ THUẬT (GOLDEN VIVA CHEAT-SHEET)")
    lines.append("")
    lines.append("| Câu hỏi của Thầy | Thành viên | Câu trả lời vàng (3 giây đầu ghi điểm 10) |")
    lines.append("| :--- | :---: | :--- |")
    golden_answers = [
        ("Token đăng nhập lưu ở đâu?", "TV2 (Giang)", "*Lưu trong chip mã hóa phần cứng SecureStore (Keychain/Keystore), tuyệt đối cấm AsyncStorage.*"),
        ("Dữ liệu bài tập lưu ở đâu?", "TV2 (Giang)", "*Lưu trong bộ nhớ đệm RAM của TanStack Query v5 với staleTime 1 phút và retry 2 lần.*"),
        ("Mất mạng dữ liệu lưu ở đâu?", "Cả 2", "*Lưu trong AsyncStorage kèm mốc thời gian thực tế hiển thị trên Banner hổ phách.*"),
        ("Optimistic Update hoạt động ra sao?", "TV2 (Giang)", "*onMutate chụp snapshot cache cũ -> onError rollback -> onSettled đồng bộ server.*"),
        ("Tiến độ % lưu ở đâu trong DB?", "TV2 (Giang)", "*Không lưu trong DB! Là Derived State O(N) tính toán động để tránh Data Inconsistency.*"),
        ("Ảnh đính kèm lưu ở đâu?", "TV1 (Trâm)", "*Lưu đường dẫn Local URI file:/// trong sandbox app, không lưu chuỗi Base64 tránh nghẽn RAM.*"),
        ("Tại sao không lỗi múi giờ 48h?", "TV2 (Giang)", "*Dùng hàm startOfDay quy chuẩn về 00:00:00 ngày lịch, không cộng trừ mili-giây chay.*"),
        ("Deep Link hoạt động thế nào?", "TV1 (Trâm)", "*Payload notification chứa assignmentId, Listener ở Root Layout đón nhận điều hướng sang [id].*"),
        ("Luật Zero-Hex hoạt động ra sao?", "TV1 (Trâm)", "*Mã hex tập trung duy nhất ở tokens.ts, UI dùng Semantic Tokens đổi Dark Mode tự động.*"),
        ("Vùng chạm 44x44pt ở đâu?", "TV1 (Trâm)", "*Cố định minHeight/minWidth 44pt tại AccessibleButton và ô checkbox, đạt chuẩn Apple HIG.*"),
        ("StateHandler làm nhiệm vụ gì?", "TV2 (Giang)", "*Xử lý đủ 4 trạng thái: Loading, Empty, Error có nút Retry gọi refetch thật, và Content.*"),
        ("Tại sao phân tách 4 vùng state?", "Cả 2", "*Áp dụng Separation of Concerns: Server State (TanStack), Client State (Zustand), Offline (AsyncStorage), Secure (SecureStore).*")
    ]
    for q, m, a in golden_answers:
        lines.append(f"| {q} | {m} | {a} |")
    lines.append("")
    lines.append("---")
    lines.append("*Hồ sơ chuyên đề kiến trúc & dữ liệu được biên soạn phục vụ đồ án Mobile Final Project — DNTU 2026.*")
    return "\n".join(lines)

if __name__ == "__main__":
    base_dir = r"c:\CaNhan\LapTrinhMobile\1923050181_1923051154_StudyPlanner"
    output_docx = os.path.join(base_dir, "Thuyet_Trinh_Chuyen_Sau_Kien_Truc_Du_An.docx")
    output_md = os.path.join(base_dir, "THUYET_TRINH_KIEN_TRUC_VA_DU_LIEU.md")
    
    build_architecture_documents(output_docx, output_md)
    
    # Also copy to parent
    parent_docx = r"c:\CaNhan\LapTrinhMobile\Thuyet_Trinh_Chuyen_Sau_Kien_Truc_Du_An.docx"
    try:
        shutil.copyfile(output_docx, parent_docx)
        print(f"Copied to parent folder: {parent_docx}")
    except Exception as e:
        print(f"Notice on copy to parent: {e}")
